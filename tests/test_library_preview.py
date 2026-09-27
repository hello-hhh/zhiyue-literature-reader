"""Tests for Word/web preview text and citation locations."""

from docx import Document

from src.library.preview import (
    citation_location,
    extract_docx_text,
    extract_pdf_page_text,
    get_pdf_page_count,
    text_to_html,
)


def test_extract_docx_text_keeps_paragraphs_and_tables(tmp_path):
    path = tmp_path / "demo.docx"
    document = Document()
    document.add_paragraph("研究背景")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "方法"
    table.cell(0, 1).text = "实验"
    document.save(path)

    text = extract_docx_text(path)
    assert "研究背景" in text
    assert "方法 | 实验" in text


def test_text_to_html_escapes_content_and_marks_excerpt():
    rendered = text_to_html("安全 <script>\n需要定位", "需要定位")
    assert "<script>" not in rendered
    assert "&lt;script&gt;" in rendered
    assert "<mark>需要定位</mark>" in rendered


def test_pdf_citation_returns_page_number():
    page, excerpt = citation_location(
        {"type": "pdf", "start_line": 7, "text": "引用内容"}
    )
    assert page == 7
    assert excerpt == "引用内容"


def test_pdf_page_text_can_be_extracted(tmp_path):
    import pymupdf

    path = tmp_path / "demo.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Legal research page")
    document.save(path)
    document.close()

    assert get_pdf_page_count(path) == 1
    assert "Legal research page" in extract_pdf_page_text(path, 1)
