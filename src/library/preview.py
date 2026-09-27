"""Document preview extraction used by the two-column reading workspace."""

from __future__ import annotations

import html
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document


def extract_docx_text(path: str | Path) -> str:
    document = Document(str(path))
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                lines.append(" | ".join(values))
    return "\n\n".join(lines)


def get_pdf_page_count(path: str | Path) -> int:
    """Return the number of pages in a PDF without keeping the file open."""
    import pymupdf

    with pymupdf.open(str(path)) as document:
        return document.page_count


def extract_pdf_page_text(path: str | Path, page: int) -> str:
    """Extract one 1-based PDF page for the highlightable reading view."""
    import pymupdf

    with pymupdf.open(str(path)) as document:
        if document.page_count == 0:
            return ""
        page_index = min(max(int(page), 1), document.page_count) - 1
        return document[page_index].get_text("text").strip()


def load_preview_text(document: dict) -> str:
    """Return readable text for Word, webpage, and plain-text documents."""
    cached = (document.get("preview_text") or "").strip()
    if cached:
        return cached
    path = Path(document["stored_path"])
    source_type = document.get("source_type", "")
    if source_type == "docx" and path.suffix.lower() == ".docx":
        return extract_docx_text(path)
    if source_type == "url" or path.suffix.lower() in {".html", ".htm"}:
        raw = path.read_text(encoding="utf-8", errors="replace")
        return BeautifulSoup(raw, "lxml").get_text("\n", strip=True)
    if path.suffix.lower() in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="replace")
    return "当前格式已完成索引，但暂不支持在页面内还原版式。"


def text_to_html(text: str, highlight: str = "") -> str:
    """Convert plain extracted text to safe, reader-friendly HTML."""
    safe = html.escape(text)
    if highlight:
        escaped_highlight = html.escape(highlight[:240])
        safe = safe.replace(escaped_highlight, f"<mark>{escaped_highlight}</mark>", 1)
    paragraphs = [part.strip() for part in safe.split("\n") if part.strip()]
    return "".join(f"<p>{part}</p>" for part in paragraphs)


def citation_location(entry: dict) -> tuple[int, str]:
    """Return the preview page and excerpt for a retrieved chunk."""
    page = int(entry.get("start_line", 1) or 1) if entry.get("type") == "pdf" else 1
    return page, str(entry.get("text", ""))
